"""DOCX paragraph and table extraction using python-docx."""
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.parsers.base import BaseParser, DocumentParseError, ensure_text


class DOCXParser(BaseParser):
    def parse(self, path: Path) -> dict:
        try:
            document = Document(str(path))
            blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            text = ensure_text("\n".join(blocks), path.name)
            return {"text": text, "pages": [{"page_number": None, "text": text}],
                    "metadata": {"format": "docx", "page_count": None,
                                 "paragraph_count": len(document.paragraphs), "table_count": len(document.tables)}}
        except DocumentParseError:
            raise
        except (PackageNotFoundError, BadZipFile, OSError, ValueError, KeyError) as exc:
            raise DocumentParseError(f"Could not parse DOCX '{path.name}': {exc}") from exc
